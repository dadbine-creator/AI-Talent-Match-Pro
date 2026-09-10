# cv_extract.py
# ============================================================
# AI Talent Match Pro — CV text extraction (Team DNA)
#
# Turns an uploaded PDF / DOCX / TXT into plain text. Deliberately small:
# extraction failures are returned as values, never raised, so one corrupt
# file in a 200-file bulk upload cannot take down the whole batch.
#
# The only new third-party dependencies in the whole Team DNA feature are
# the two importable here, and both are optional at import time so the app
# still boots if a deploy misses them.
# ============================================================
from __future__ import annotations

import io
import logging
import os
from typing import Optional, Tuple

logger = logging.getLogger(__name__)

# Optional imports — a missing extractor degrades that ONE format, not the app.
try:
    import pdfplumber  # type: ignore
except Exception:                                     # pragma: no cover - env dependent
    pdfplumber = None

try:
    import docx  # python-docx                        # type: ignore
except Exception:                                     # pragma: no cover - env dependent
    docx = None


# Hard ceilings. A CV is a few KB of text; anything past this is a paste bomb
# or a scanned book, and it would blow the model's context either way.
MAX_FILE_BYTES = 10 * 1024 * 1024      # 10 MB per file
MAX_TEXT_CHARS = 30_000                # matches MAX_RESUME_CHARS in main.py
MIN_USEFUL_CHARS = 40                  # below this there is nothing to score

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".text", ".md")


class ExtractionError(Exception):
    """Raised only by _extract_* helpers; extract_text() converts it to a value."""


def _ext(filename: str) -> str:
    return os.path.splitext(filename or "")[1].lower()


def _extract_pdf(data: bytes) -> str:
    if pdfplumber is None:
        raise ExtractionError("PDF support is not installed on the server (pdfplumber missing)")
    try:
        chunks = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                chunks.append(page.extract_text() or "")
        return "\n".join(chunks)
    except ExtractionError:
        raise
    except Exception as e:
        # Encrypted, truncated, or not actually a PDF.
        raise ExtractionError(f"Could not read this PDF ({type(e).__name__})")


def _extract_docx(data: bytes) -> str:
    if docx is None:
        raise ExtractionError("DOCX support is not installed on the server (python-docx missing)")
    try:
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        # Tables carry real content in a lot of CV templates.
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)
    except ExtractionError:
        raise
    except Exception as e:
        raise ExtractionError(f"Could not read this DOCX ({type(e).__name__})")


def _extract_txt(data: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    raise ExtractionError("Could not decode this text file")


def extract_text(filename: str, data: bytes) -> Tuple[Optional[str], Optional[str]]:
    """Extract plain text from one uploaded file.

    Returns (text, error). Exactly one is non-None. Never raises — callers
    processing a batch rely on that to keep going past a bad file.
    """
    if not data:
        return None, "File is empty"
    if len(data) > MAX_FILE_BYTES:
        return None, f"File is larger than {MAX_FILE_BYTES // (1024 * 1024)}MB"

    extension = _ext(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        return None, f"Unsupported file type '{extension or 'unknown'}' — use PDF, DOCX or TXT"

    try:
        if extension == ".pdf":
            text = _extract_pdf(data)
        elif extension == ".docx":
            text = _extract_docx(data)
        else:
            text = _extract_txt(data)
    except ExtractionError as e:
        return None, str(e)
    except Exception as e:                             # pragma: no cover - defensive
        logger.warning("cv_extract: unexpected failure on %s: %s", filename, e)
        return None, f"Could not read this file ({type(e).__name__})"

    text = _normalize(text)
    if len(text) < MIN_USEFUL_CHARS:
        # The common cause is a scanned/image-only PDF with no text layer.
        return None, "No readable text found — if this is a scanned document it needs OCR first"
    return text[:MAX_TEXT_CHARS], None


def _normalize(text: str) -> str:
    """Collapse the whitespace soup that PDF extraction produces."""
    if not text:
        return ""
    lines = [line.strip() for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    out, blank = [], 0
    for line in lines:
        if line:
            out.append(line)
            blank = 0
        else:
            blank += 1
            if blank <= 1:          # keep single blank lines, drop runs
                out.append("")
    return "\n".join(out).strip()


def guess_name(filename: str, text: str) -> str:
    """Best-effort candidate name for display.

    Prefers the first plausible line of the CV, falls back to the filename.
    Never invents a name — worst case you get the filename back.
    """
    for line in (text or "").split("\n")[:5]:
        line = line.strip()
        # A name line: short, no digits, no @, 1-5 words.
        if (
            2 <= len(line) <= 60
            and "@" not in line
            and not any(ch.isdigit() for ch in line)
            and 1 <= len(line.split()) <= 5
            and not line.lower().startswith(("curriculum", "resume", "cv", "profile"))
        ):
            return line
    stem = os.path.splitext(os.path.basename(filename or ""))[0]
    stem = stem.replace("_", " ").replace("-", " ").strip()
    return stem[:120] or "Unnamed candidate"
